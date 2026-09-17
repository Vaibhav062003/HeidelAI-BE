import asyncio
import io
import json
import os
from datetime import date, timedelta
from pathlib import Path

os.environ.setdefault("MONGODB_URI", "mongodb://127.0.0.1:27017")
os.environ.setdefault("MONGODB_CLUSTER", "native_invoice_reporting_test")

from fastapi import HTTPException

import services.native_invoice_service as invoice_service
from services.native_invoice_service import csv_cell, export_invoices_json, export_invoices_xlsx, reporting_period


class _Cursor:
    def __init__(self, rows):
        self.rows = rows

    def sort(self, *_args):
        return self

    async def to_list(self, **_kwargs):
        return self.rows


class _Invoices:
    def __init__(self, rows):
        self.rows = rows

    def find(self, *_args, **_kwargs):
        return _Cursor(self.rows)


async def _export_contract() -> None:
    original = invoice_service.invoices
    invoice_service.invoices = _Invoices([{
        "invoice_id": "invoice-1", "invoice_date": "2026-09-17", "due_date": None,
        "lifecycle": "issued", "revision": 2, "customer": {"name": "=2+2", "gstin": ""},
        "items": [], "total_paise": 10_000, "taxable_value_paise": 10_000,
        "amount_paid_paise": 0, "amount_refunded_paise": 0, "payments": [], "refunds": [],
        "pdf": {"public_id": "private"}, "receipt": {"public_id": "private-receipt"},
        "signature_image": {"filename": "signature.png", "content_type": "image/png", "byte_size": 1, "sha256": "a" * 64, "data": b"x"},
    }])
    try:
        data, filename = await export_invoices_json("org-a", "2026-09-17", "2026-09-17")
        exported = json.loads(data)
        assert filename.endswith(".json")
        assert "pdf" not in exported["invoices"][0] and "receipt" not in exported["invoices"][0]
        assert "data" not in exported["invoices"][0]["signature_image"]

        data, filename = await export_invoices_xlsx("org-a", "2026-09-17", "2026-09-17")
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=False)["Invoices"]
        assert filename.endswith(".xlsx")
        assert sheet.cell(2, 6).value == "'=2+2"
    finally:
        invoice_service.invoices = original


def main() -> None:
    today = date.fromisoformat(invoice_service._today())
    assert reporting_period(today, today) == (today.isoformat(), today.isoformat())
    assert csv_cell("Customer") == "Customer"
    for value in ("=SUM(A1:A2)", "+cmd", "-1+2", "@formula", "  =hidden"):
        assert csv_cell(value).lstrip().startswith("'")
    for start, end in ((today + timedelta(days=1), today), (today - timedelta(days=732), today), (today, today + timedelta(days=1))):
        try:
            reporting_period(start, end)
        except HTTPException as exc:
            assert exc.status_code == 422
        else:
            raise AssertionError("Unsafe report date range was accepted")
    service = Path("services/native_invoice_service.py").read_text()
    assert '"net_received_paise": 1' in service
    assert '"collected_paise"' not in service
    asyncio.run(_export_contract())


if __name__ == "__main__":
    main()
    print("Native invoice reporting boundaries passed.")
